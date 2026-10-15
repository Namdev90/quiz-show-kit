#!/usr/bin/env python3
"""Offline DOCX bank interchange. Export uses installed python-docx; import uses bounded stdlib XML parsing. MIT."""
import base64, io, json, os, re, sys, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R='{http://schemas.openxmlformats.org/package/2006/relationships}'
META='customXml/qsk-bank.xml'
FIELDS=[('Status','status'),('Round','round'),('Pair ID','pairId'),('Topic','topic'),('Difficulty','difficulty'),('Question type','subtype'),('Question','stem'),('Correct option ID','correctOptionId'),('Answer explanation','explanation')]
SOURCE_FIELDS=[('Title','title'),('URL','url'),('Section or page','section'),('Edition','edition'),('Publication date','publicationDate'),('Access date','accessedAt'),('Verification date','verifiedAt'),('Status','status'),('Support note','supportNote')]

def export(source,destination):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 data=json.loads(Path(source).read_text(encoding='utf8'));doc=Document()
 section=doc.sections[0];section.top_margin=section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.7)
 for name in ['Normal','Title','Subtitle','Caption','Heading 1','Heading 2']:
  style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
 for border in doc.styles['Title'].element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
 doc.styles['Normal'].font.size=Pt(10)
 doc.styles['Normal'].paragraph_format.space_after=Pt(3)
 doc.add_paragraph('Editable question bank',style='Title')
 doc.add_paragraph(data.get('eventTitle','Quiz Show Kit'),style='Subtitle')
 doc.add_paragraph('QSK1 event '+data['eventId']+' revision '+data['baseFingerprint'],style='Caption')
 doc.add_paragraph('Edit the right hand cells, then return this DOCX to your quiz assistant for a change preview. Your current quiz is not changed until you explicitly apply the preview. This document includes answer keys and is for the creator or host.')
 doc.add_paragraph('Keep Question ID and Option IDs stable. To change the key, copy the chosen Option ID into Correct option ID. Keep four options. Copy a complete question table to add an item with a new Question ID and four new Option IDs. Delete a whole table to remove an item. Main and tie questions must remain matched pairs.')
 doc.add_paragraph('Question types: text, numerical, image, graph, code, sequence or matching. Every type uses selectable MCQ answers. Media and narration references are preserved from the exported bank; supply approved replacement media separately. Source dates use YYYY-MM-DD. Empty source records are ignored. Sources and explanations require independent checking after any educational change.')
 doc.add_paragraph('If a field is confusing, ask “I do not understand — fix it for me”. The assistant can repair harmless formatting and show a preview. It must ask before changing meaning or resolving an answer disagreement.')
 for qi,q in enumerate(data['questions']):
  doc.add_page_break();doc.add_paragraph('Question '+q['id'],style='Heading 1')
  rows=[('Question ID',q['id'])]+[(label,str(q.get(key,''))) for label,key in FIELDS]
  for i,label in enumerate('ABCD'):
   rows += [('Option '+label+' ID',q['optionIds'][i]),('Option '+label,q['options'][i]),('Explanation '+label,q['optionExplanations'][i])]
  rows += [('Media reference',q.get('visual',{}).get('key',q.get('visual',{}).get('type',''))),('Media alternative text',q.get('visual',{}).get('alt',''))]
  v=q.get('verification',{});rows += [('Verification status',v.get('status','unverified')),('Checked date',v.get('checkedAt','')),('Verification note',v.get('note','')),('Override reason',v.get('override',{}).get('reason','')),('Override chosen option ID',v.get('override',{}).get('chosenOptionId','')),('Override recorded date',v.get('override',{}).get('recordedAt',''))]
  for i,s in enumerate(q.get('sources',[]) or [{}],1):
   rows += [('Source '+str(i)+' '+label,str(s.get(key,''))) for label,key in SOURCE_FIELDS]
  table=doc.add_table(rows=0,cols=2);table.style='Table Grid';table.autofit=False;table.columns[0].width=Inches(1.75);table.columns[1].width=Inches(5.3)
  for label,value in rows:
   cells=table.add_row().cells;cells[0].text=label;cells[1].text=value
   cells[0].width=Inches(1.75);cells[1].width=Inches(5.3)
   trPr=table.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
   for cell in cells:
    tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tcPr.append(borders)
   for run in cells[0].paragraphs[0].runs:run.bold=True
  doc.add_paragraph('Verification and source status are recorded author decisions. A link alone does not establish verification. Creator overrides retain the unresolved disagreement.',style='Caption')
 buffer=io.BytesIO();doc.save(buffer)
 bank=ET.Element('{urn:quiz-show-kit:bank:1}bank');bank.text=json.dumps(data,ensure_ascii=False,separators=(',',':'))
 with zipfile.ZipFile(buffer) as old,zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as output:
  for item in old.infolist():
   content=old.read(item.filename)
   if item.filename=='word/_rels/document.xml.rels':
    ET.register_namespace('',R[1:-1]);rels=ET.fromstring(content);ET.SubElement(rels,R+'Relationship',Id='qskQuestionBank',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml',Target='../customXml/qsk-bank.xml');content=ET.tostring(rels,encoding='utf-8',xml_declaration=True)
   output.writestr(item,content)
  output.writestr(META,ET.tostring(bank,encoding='utf-8',xml_declaration=True))
 return {'questions':len(data['questions']),'output':str(destination)}

def xml(content,label):
 if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():raise ValueError(label+': document type declarations and entities are unsupported')
 root=ET.fromstring(content)
 count=0;stack=[(root,0)]
 while stack:
  node,depth=stack.pop();count+=1
  if depth>45 or count>100000:raise ValueError(label+': XML nesting or size limit exceeded')
  stack.extend((child,depth+1) for child in node)
 return root

def load(source,fallback=None):
 if os.path.getsize(source)>8*1024*1024:raise ValueError('Document exceeds the 8 MB import limit')
 with zipfile.ZipFile(source) as z:
  info=z.infolist();names=[x.filename for x in info]
  if len(info)>300 or len(names)!=len(set(names)) or sum(x.file_size for x in info)>15*1024*1024:raise ValueError('Document has duplicate parts or exceeds the expanded import limit')
  if any(x.flag_bits&1 or x.file_size>5*1024*1024 for x in info):raise ValueError('Encrypted or oversized document parts are unsupported')
  for name in names:
   if name.startswith('/') or '..' in name.replace('\\','/').split('/') or re.search(r'(vbaProject|embeddings/|activeX/)',name,re.I):raise ValueError('Macros, embedded objects or unsafe package paths are unsupported')
  for name in names:
   if name.endswith('.rels'):
    root=xml(z.read(name),name)
    for rel in root:
     if rel.get('TargetMode')=='External' and (not rel.get('Type','').endswith('/hyperlink') or not re.match(r'^https://[^\s<>]+$',rel.get('Target',''))):raise ValueError('External document resources are unsupported; HTTPS source hyperlinks are never fetched')
  root=xml(z.read('word/document.xml'),'Question document')
  if META in names:metadata=json.loads(xml(z.read(META),'Recovery metadata').text or '{}')
  else:
   headers=[''.join(n.text or '' for n in p.iter(W+'t')) for p in root.findall('.//'+W+'body/'+W+'p')]
   identity=[re.fullmatch(r'QSK1 event ([a-zA-Z0-9._-]+) revision ([0-9a-f]{8})',x) for x in headers];identity=[x for x in identity if x]
   if not fallback or len(identity)!=1:raise ValueError('Recovery metadata and template identity are missing. Export a fresh template; your current bank is unchanged')
   metadata=json.loads(Path(fallback).read_text(encoding='utf8'))
   if identity[0].group(1)!=metadata.get('eventId') or identity[0].group(2)!=metadata.get('baseFingerprint'):raise ValueError('The template identity does not match the current bank. Reconcile with a fresh export before applying')
  if metadata.get('format')!='QSK_DOCX_1' or not isinstance(metadata.get('questions'),list):raise ValueError('Unsupported question bank template')
  if any(root.iter(W+'altChunk')) or any(root.iter(W+'object')) or any(root.iter(W+'drawing')) or any(root.iter(W+'fldSimple')) or any(root.iter(W+'instrText')) or any(root.iter(W+'ins')) or any(root.iter(W+'del')):raise ValueError('Embedded media, fields, tracked changes or unsupported content found. Accept or reject tracked changes and use media references, then preview again')
  tables=root.findall('.//'+W+'body/'+W+'tbl')
  if len(tables)>250:raise ValueError('At most 250 question tables are supported')
  existing={q['id']:q for q in metadata['questions']};questions=[];repairs=[]
  def celltext(cell):return '\n'.join(''.join(node.text or '' for node in p.iter(W+'t')) for p in cell.findall(W+'p')).strip()
  for table in tables:
   fields={}
   for row in table.findall(W+'tr'):
    cells=row.findall(W+'tc')
    if len(cells)!=2 or any(c.findall('.//'+W+'tbl') for c in cells):raise ValueError('A question table needs exactly two columns and no nested tables')
    label,value=map(celltext,cells)
    if label in fields:raise ValueError('Question '+fields.get('Question ID','unknown')+': duplicate field '+label)
    fields[label]=value
   qid=fields.get('Question ID','');context='Question '+(qid or 'unknown')
   required=['Question ID']+[x[0] for x in FIELDS]+[prefix+' '+letter for letter in 'ABCD' for prefix in ['Option','Explanation']]+['Option '+x+' ID' for x in 'ABCD']
   for key in required:
    if key not in fields:raise ValueError(context+': missing field '+key+'. Copy its row from a fresh template')
   known=set(required+['Media reference','Media alternative text','Verification status','Checked date','Verification note','Override reason','Override chosen option ID','Override recorded date'])
   for key in fields:
    if key not in known and not any(re.fullmatch(r'Source (?:[1-9]|1[0-2]) '+re.escape(label),key) for label,_ in SOURCE_FIELDS):raise ValueError(context+': unsupported field '+key)
   q=json.loads(json.dumps(existing.get(qid,{})));q['id']=qid
   for label,key in FIELDS:q[key]=fields[label]
   if re.fullmatch(r'\d+',str(q['round'])):q['round']=int(q['round'])
   q['options']=[fields['Option '+x] for x in 'ABCD'];q['optionIds']=[fields['Option '+x+' ID'] for x in 'ABCD'];q['optionExplanations']=[fields['Explanation '+x] for x in 'ABCD'];q.pop('correct',None)
   media=fields.get('Media reference','')
   if media:
    q.setdefault('visual',{});q['visual']['key']=media;q['visual']['alt']=fields.get('Media alternative text','')
   else:q.pop('visual',None)
   # Original absent verification metadata remains absent on a truly unchanged export/import.
   v={'status':fields.get('Verification status','unverified')}
   for label,key in [('Checked date','checkedAt'),('Verification note','note')]:
    if fields.get(label):v[key]=fields[label]
   if any(fields.get(x) for x in ['Override reason','Override chosen option ID','Override recorded date']):
    v['override']={'reason':fields.get('Override reason',''),'chosenOptionId':fields.get('Override chosen option ID','')}
    if fields.get('Override recorded date'):v['override']['recordedAt']=fields['Override recorded date']
   if 'verification' in q or v!={'status':'unverified'}:q['verification']=v
   sources=[]
   for i in range(1,13):
    record={key:fields['Source '+str(i)+' '+label] for label,key in SOURCE_FIELDS if fields.get('Source '+str(i)+' '+label)}
    if record:sources.append(record)
   q['sources']=sources;questions.append(q)
  return {'eventId':metadata.get('eventId'),'baseFingerprint':metadata.get('baseFingerprint'),'questions':questions,'repairs':repairs}

if __name__=='__main__':
 try:
  command=sys.argv[1]
  if command=='export':result=export(sys.argv[2],sys.argv[3])
  elif command=='import':result=load(sys.argv[2],sys.argv[3] if len(sys.argv)>3 else None)
  else:raise ValueError('Choose export or import')
  print(json.dumps(result,ensure_ascii=False))
 except Exception as exc:
  print(json.dumps({'error':str(exc)},ensure_ascii=False),file=sys.stderr);sys.exit(1)
